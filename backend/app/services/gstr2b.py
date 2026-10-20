"""GSTR-2B reconciliation: does what suppliers reported match the approved books?

GSTR-2B is the government's auto-drafted statement of the input tax credit (ITC)
that suppliers reported for a month. For each supplier invoice in it:

    matched     same supplier GSTIN and invoice number; date and amounts agree
    mismatched  same supplier and number, but the date or an amount differs
    possible    same supplier and amounts, but the number is written differently
    only_2b     the supplier reported it; it is not in the approved books

and, from the books' side, approved invoices no statement matched are
"in books only": ITC at risk until the supplier reports them.

FILE FORMATS:
- Excel (.xlsx), as downloaded from the GST portal: checked against a real file
  (column layout and value formats only; its contents are never stored in the repo).
  Invoices come from the "B2B" sheet, the period from the "Read me" sheet.
- JSON: follows the commonly documented layout (b2b -> ctin/trdnm -> inv: inum,
  dt, val, pos, rev, itcavl -> items: rt, txval, igst, cgst, sgst, cess), written
  from public descriptions and not yet checked against a portal file.
Only B2B invoices are read: amendments, credit/debit notes, ISD and imports are not.
"""

import re
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Gstr2bInvoice, Gstr2bStatement, Invoice, InvoiceStatus, MatchStatus

FORMAT_NAME = "gstr2b-json (documented layout)"
EXCEL_FORMAT = "gstr2b-excel (GST portal)"
# Formats whose layout was checked against a real file downloaded from the portal.
VERIFIED_FORMATS = {EXCEL_FORMAT}
TOLERANCE = Decimal("1.00")  # rupees: rounding on either side
POSSIBLE_DATE_WINDOW = timedelta(days=5)


class Gstr2bFormatError(ValueError):
    """The file isn't a GSTR-2B JSON we can read; the message says where."""


@dataclass
class Row2B:
    supplier_gstin: str
    invoice_number: str
    supplier_name: str | None = None
    invoice_date: date | None = None
    invoice_value: Decimal | None = None
    place_of_supply: str | None = None
    reverse_charge: bool = False
    itc_available: bool = True
    taxable_value: Decimal = Decimal(0)
    igst: Decimal = Decimal(0)
    cgst: Decimal = Decimal(0)
    sgst: Decimal = Decimal(0)
    cess: Decimal = Decimal(0)


@dataclass
class Statement2B:
    period: date  # first day of the return month
    recipient_gstin: str | None
    rows: list[Row2B] = field(default_factory=list)
    source_format: str = FORMAT_NAME


# ---------------------------------------------------------------------------
# Parsing
# ---------------------------------------------------------------------------


def _amount(value, where: str) -> Decimal:
    if value in (None, ""):
        return Decimal(0)
    try:
        return Decimal(str(value)).quantize(Decimal("0.01"))
    except InvalidOperation:
        raise Gstr2bFormatError(f"{where}: not an amount: {value!r}") from None


def _date(value, where: str) -> date | None:
    if not value:
        return None
    for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(str(value), fmt).date()
        except ValueError:
            continue
    raise Gstr2bFormatError(f"{where}: not a date: {value!r}")


def _period(value) -> date:
    """'072026' (MMYYYY, as in the return period) or '2026-07' -> 2026-07-01."""
    text = str(value or "").strip()
    match = re.fullmatch(r"(\d{2})(\d{4})", text) or re.fullmatch(r"(\d{4})-(\d{2})", text)
    if not match:
        raise Gstr2bFormatError(f"return period: expected MMYYYY, got {value!r}")
    a, b = match.groups()
    year, month = (int(b), int(a)) if len(a) == 2 else (int(a), int(b))
    if not 1 <= month <= 12:
        raise Gstr2bFormatError(f"return period: no month {month}")
    return date(year, month, 1)


def _yes(value) -> bool:
    return str(value or "").strip().upper() in ("Y", "YES", "TRUE", "1")


def parse(data: dict) -> Statement2B:
    """Read a GSTR-2B JSON (documented layout) into a Statement2B."""
    if not isinstance(data, dict):
        raise Gstr2bFormatError("expected a JSON object")
    root = data.get("data", data)
    period = _period(root.get("rtnprd"))
    docdata = root.get("docdata")
    if not isinstance(docdata, dict):
        raise Gstr2bFormatError("no 'docdata' section (is this a GSTR-2B file?)")
    statement = Statement2B(period=period, recipient_gstin=root.get("gstin"))
    for s, supplier in enumerate(docdata.get("b2b") or []):
        gstin = str(supplier.get("ctin") or "").strip().upper()
        if not gstin:
            raise Gstr2bFormatError(f"b2b[{s}]: no supplier GSTIN ('ctin')")
        for i, inv in enumerate(supplier.get("inv") or []):
            where = f"b2b[{s}].inv[{i}]"
            number = str(inv.get("inum") or "").strip()
            if not number:
                raise Gstr2bFormatError(f"{where}: no invoice number ('inum')")
            row = Row2B(
                supplier_gstin=gstin,
                supplier_name=supplier.get("trdnm"),
                invoice_number=number,
                invoice_date=_date(inv.get("dt"), f"{where}.dt"),
                invoice_value=_amount(inv.get("val"), f"{where}.val"),
                place_of_supply=(str(inv.get("pos")).zfill(2) if inv.get("pos") else None),
                reverse_charge=_yes(inv.get("rev")),
                itc_available=inv.get("itcavl") is None or _yes(inv.get("itcavl")),
            )
            for j, item in enumerate(inv.get("items") or inv.get("itms") or []):
                item = item.get("itm_det", item)  # some layouts nest the amounts
                at = f"{where}.items[{j}]"
                row.taxable_value += _amount(item.get("txval"), f"{at}.txval")
                row.igst += _amount(item.get("igst", item.get("iamt")), f"{at}.igst")
                row.cgst += _amount(item.get("cgst", item.get("camt")), f"{at}.cgst")
                row.sgst += _amount(item.get("sgst", item.get("samt")), f"{at}.sgst")
                row.cess += _amount(item.get("cess", item.get("csamt")), f"{at}.cess")
            statement.rows.append(row)
    return statement


MONTHS = ["january", "february", "march", "april", "may", "june", "july", "august",
          "september", "october", "november", "december"]  # fmt: skip

# B2B sheet columns we read, by their (normalised) header text in the portal's file.
EXCEL_COLUMNS = {
    "gstinofsupplier": "supplier_gstin",
    "tradelegalname": "supplier_name",
    "invoicenumber": "invoice_number",
    "invoicedate": "invoice_date",
    "invoicevalue": "invoice_value",
    "placeofsupply": "place_of_supply",
    "supplyattractreversecharge": "reverse_charge",
    "taxablevalue": "taxable_value",
    "integratedtax": "igst",
    "centraltax": "cgst",
    "stateuttax": "sgst",
    "cess": "cess",
    "itcavailability": "itc_available",
}


def _header_key(text) -> str:
    """'Invoice Value(₹)' -> 'invoicevalue', 'State/UT Tax(₹)' -> 'stateuttax'."""
    return re.sub(r"[^a-z]", "", re.sub(r"\(.*?\)", "", str(text or "").lower()))


def _excel_period(workbook) -> date:
    """'Tax Period' (a month name) and 'Financial Year' ('2026-27') from the Read me sheet."""
    if "Read me" not in workbook.sheetnames:
        raise Gstr2bFormatError("no 'Read me' sheet (is this the portal's GSTR-2B Excel?)")
    labels = {}
    for row in workbook["Read me"].iter_rows(min_row=1, max_row=20, values_only=True):
        if row and row[0]:
            value = next((c for c in row[1:] if c not in (None, "")), None)
            labels[_header_key(row[0])] = value
    month_name = str(labels.get("taxperiod") or "").strip().lower()
    year = re.match(r"(\d{4})-(\d{2})", str(labels.get("financialyear") or ""))
    if month_name not in MONTHS or not year:
        raise Gstr2bFormatError("Read me sheet: no 'Tax Period' month or 'Financial Year'")
    month = MONTHS.index(month_name) + 1
    start_year = int(year.group(1))
    return date(start_year if month >= 4 else start_year + 1, month, 1)  # April-March year


def parse_excel(content: bytes) -> Statement2B:
    """Read the GST portal's GSTR-2B Excel download (B2B sheet) into a Statement2B."""
    import io

    import openpyxl

    from app.services.gst import state_code_of_place

    try:
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise Gstr2bFormatError("not an Excel file we can open") from None
    period = _excel_period(workbook)
    recipient = None
    for row in workbook["Read me"].iter_rows(min_row=1, max_row=20, values_only=True):
        if row and _header_key(row[0]) == "gstin":
            recipient = next((str(c).strip() for c in row[1:] if c not in (None, "")), None)
    if "B2B" not in workbook.sheetnames:
        raise Gstr2bFormatError("no 'B2B' sheet")
    rows = list(workbook["B2B"].iter_rows(values_only=True))
    header_at = next(
        (i for i, row in enumerate(rows) if row and _header_key(row[0]) == "gstinofsupplier"),
        None,
    )
    if header_at is None or header_at + 1 >= len(rows):
        raise Gstr2bFormatError("B2B sheet: no 'GSTIN of supplier' header")
    # Two header rows: group titles ("Invoice Details", "Tax Amount") above sub-columns.
    columns: dict[str, int] = {}
    top, sub = rows[header_at], rows[header_at + 1]
    for index in range(max(len(top), len(sub))):
        for text in (sub[index] if index < len(sub) else None, top[index]):
            key = _header_key(text)
            if key in EXCEL_COLUMNS and EXCEL_COLUMNS[key] not in columns:
                columns[EXCEL_COLUMNS[key]] = index
                break
    missing = {"supplier_gstin", "invoice_number", "invoice_date", "taxable_value"} - set(columns)
    if missing:
        raise Gstr2bFormatError(f"B2B sheet: missing column(s) {sorted(missing)}")

    statement = Statement2B(period=period, recipient_gstin=recipient, source_format=EXCEL_FORMAT)
    for number, row in enumerate(rows[header_at + 2 :], start=header_at + 3):

        def cell(name, row=row):
            index = columns.get(name)
            return row[index] if index is not None and index < len(row) else None

        gstin = str(cell("supplier_gstin") or "").strip().upper()
        if not gstin:
            continue  # blank rows at the end
        where = f"B2B row {number}"
        value = cell("invoice_date")
        statement.rows.append(
            Row2B(
                supplier_gstin=gstin,
                supplier_name=str(cell("supplier_name") or "").strip() or None,
                invoice_number=str(cell("invoice_number") or "").strip(),
                invoice_date=value.date()
                if isinstance(value, datetime)
                else _date(value, f"{where} invoice date"),
                invoice_value=_amount(cell("invoice_value"), f"{where} invoice value"),
                place_of_supply=state_code_of_place(str(cell("place_of_supply") or "")),
                reverse_charge=_yes(cell("reverse_charge")),
                itc_available=cell("itc_available") is None or _yes(cell("itc_available")),
                taxable_value=_amount(cell("taxable_value"), f"{where} taxable value"),
                igst=_amount(cell("igst"), f"{where} integrated tax"),
                cgst=_amount(cell("cgst"), f"{where} central tax"),
                sgst=_amount(cell("sgst"), f"{where} state/UT tax"),
                cess=_amount(cell("cess"), f"{where} cess"),
            )
        )
    return statement


def parse_file(content: bytes, file_name: str = "") -> Statement2B:
    """A GSTR-2B upload: the portal's Excel, or JSON."""
    if content[:4] == b"PK\x03\x04" or file_name.lower().endswith(".xlsx"):
        return parse_excel(content)
    import json

    try:
        data = json.loads(content)
    except (UnicodeDecodeError, ValueError):
        raise Gstr2bFormatError("not an Excel (.xlsx) or JSON file") from None
    return parse(data)


# ---------------------------------------------------------------------------
# Matching
# ---------------------------------------------------------------------------


def number_key(number: str | None) -> str:
    """'INV/26-27/0042' and 'inv 26 27 42' -> 'INV|26|27|42'.

    Suppliers rarely type an invoice number exactly as printed: case, spaces,
    separators and leading zeros are ignored.
    """
    parts = re.findall(r"[A-Za-z]+|\d+", (number or "").upper())
    return "|".join((p.lstrip("0") or "0") if p.isdigit() else p for p in parts)


def _money(value) -> Decimal:
    return Decimal(value or 0)


def differences(invoice: Invoice, row: Gstr2bInvoice) -> list[dict]:
    """Fields where the books and the statement disagree (amounts within ₹1 agree)."""
    pairs = [
        ("invoice_date", invoice.invoice_date, row.invoice_date),
        ("taxable_value", invoice.subtotal, row.taxable_value),
        ("igst", invoice.igst, row.igst),
        ("cgst", invoice.cgst, row.cgst),
        ("sgst", invoice.sgst, row.sgst),
        ("total", invoice.total, row.invoice_value),
    ]
    found = []
    for name, books, reported in pairs:
        if name == "invoice_date":
            same = books is None or reported is None or books == reported
        else:
            same = abs(_money(books) - _money(reported)) <= TOLERANCE
        if not same:
            found.append(
                {
                    "field": name,
                    "books": None if books is None else str(books),
                    "gstr2b": None if reported is None else str(reported),
                }
            )
    return found


def _period_end(period: date) -> date:
    return (period.replace(day=28) + timedelta(days=4)).replace(day=1) - timedelta(days=1)


def _matched_elsewhere(db: Session, statement: Gstr2bStatement) -> set[int]:
    """Invoices that another statement already matched (they can't match twice)."""
    return set(
        db.scalars(
            select(Gstr2bInvoice.matched_invoice_id)
            .join(Gstr2bStatement)
            .where(
                Gstr2bStatement.user_id == statement.user_id,
                Gstr2bStatement.id != statement.id,
                Gstr2bInvoice.match_status.in_([MatchStatus.MATCHED, MatchStatus.MISMATCHED]),
                Gstr2bInvoice.matched_invoice_id.is_not(None),
            )
        )
    )


def books_for(db: Session, statement: Gstr2bStatement) -> list[Invoice]:
    """Approved invoices this statement may match: dated up to the period's end and
    not matched by another statement (suppliers often report a month late)."""
    taken = _matched_elsewhere(db, statement)
    invoices = db.scalars(
        select(Invoice).where(
            Invoice.user_id == statement.user_id,
            Invoice.status == InvoiceStatus.APPROVED,
            Invoice.invoice_date <= _period_end(statement.period),
        )
    ).all()
    return [inv for inv in invoices if inv.id not in taken]


def reconcile(db: Session, statement: Gstr2bStatement) -> None:
    """Match every row of the statement against the books. The caller commits."""
    books = books_for(db, statement)
    by_key: dict[tuple[str, str], list[Invoice]] = {}
    for inv in books:
        key = ((inv.vendor_gstin or "").upper(), number_key(inv.invoice_number))
        by_key.setdefault(key, []).append(inv)
    used: set[int] = set()

    for row in statement.invoices:
        row.match_status, row.matched_invoice_id, row.differences = MatchStatus.ONLY_2B, None, None
        candidates = [
            inv
            for inv in by_key.get((row.supplier_gstin, number_key(row.invoice_number)), [])
            if inv.id not in used
        ]
        if candidates:
            # Several approved copies can't exist (the ledger is unique), but be safe:
            # prefer the one with the fewest differences.
            best = min(candidates, key=lambda inv: len(differences(inv, row)))
            found = differences(best, row)
            row.matched_invoice_id = best.id
            row.match_status = MatchStatus.MISMATCHED if found else MatchStatus.MATCHED
            row.differences = found or None
            used.add(best.id)

    # Second pass: same supplier and amounts, number written differently.
    for row in statement.invoices:
        if row.match_status != MatchStatus.ONLY_2B:
            continue
        options = [
            inv
            for inv in books
            if inv.id not in used
            and (inv.vendor_gstin or "").upper() == row.supplier_gstin
            and abs(_money(inv.total) - _money(row.invoice_value)) <= TOLERANCE
            and abs(_money(inv.subtotal) - _money(row.taxable_value)) <= TOLERANCE
            and (
                inv.invoice_date is None
                or row.invoice_date is None
                or abs(inv.invoice_date - row.invoice_date) <= POSSIBLE_DATE_WINDOW
            )
        ]
        if len(options) == 1:  # only suggest when it's unambiguous
            row.match_status, row.matched_invoice_id = MatchStatus.POSSIBLE, options[0].id
            row.differences = [
                {
                    "field": "invoice_number",
                    "books": options[0].invoice_number,
                    "gstr2b": row.invoice_number,
                }
            ]
            used.add(options[0].id)


def save_statement(
    db: Session, user_id: int, parsed: Statement2B, *, file_name: str
) -> Gstr2bStatement:
    """Store (replacing the same month's statement) and reconcile. The caller commits."""
    old = db.scalar(
        select(Gstr2bStatement).where(
            Gstr2bStatement.user_id == user_id, Gstr2bStatement.period == parsed.period
        )
    )
    if old is not None:
        db.delete(old)
        db.flush()
    statement = Gstr2bStatement(
        user_id=user_id,
        period=parsed.period,
        recipient_gstin=parsed.recipient_gstin,
        source_format=parsed.source_format,
        file_name=file_name[:255],
        invoices=[
            Gstr2bInvoice(
                supplier_gstin=r.supplier_gstin[:15],
                supplier_name=(r.supplier_name or None) and r.supplier_name[:255],
                invoice_number=r.invoice_number[:100],
                invoice_date=r.invoice_date,
                invoice_value=r.invoice_value,
                place_of_supply=r.place_of_supply,
                reverse_charge=r.reverse_charge,
                itc_available=r.itc_available,
                taxable_value=r.taxable_value,
                igst=r.igst,
                cgst=r.cgst,
                sgst=r.sgst,
                cess=r.cess,
            )
            for r in parsed.rows
        ],
    )
    db.add(statement)
    db.flush()
    reconcile(db, statement)
    db.flush()
    return statement


# ---------------------------------------------------------------------------
# The report
# ---------------------------------------------------------------------------


def _tax(row) -> Decimal:
    return _money(row.igst) + _money(row.cgst) + _money(row.sgst)


def report(db: Session, statement: Gstr2bStatement) -> dict:
    """Everything the Reconcile page shows, as plain data."""
    rows = statement.invoices
    matched_ids = {r.matched_invoice_id for r in rows if r.matched_invoice_id}
    start, end = statement.period, _period_end(statement.period)
    unmatched = [inv for inv in books_for(db, statement) if inv.id not in matched_ids]
    at_risk = [inv for inv in unmatched if inv.invoice_date and inv.invoice_date >= start]
    waiting = [inv for inv in unmatched if not inv.invoice_date or inv.invoice_date < start]
    counts = Counter(row.match_status for row in rows)
    claimable = sum(
        (_tax(r) for r in rows if r.match_status == MatchStatus.MATCHED and r.itc_available),
        Decimal(0),
    )
    return {
        "statement": statement,
        "period_end": end,
        "counts": {
            MatchStatus.MATCHED: counts.get(MatchStatus.MATCHED, 0),
            MatchStatus.MISMATCHED: counts.get(MatchStatus.MISMATCHED, 0),
            MatchStatus.POSSIBLE: counts.get(MatchStatus.POSSIBLE, 0),
            MatchStatus.ONLY_2B: counts.get(MatchStatus.ONLY_2B, 0),
            "books_only": len(at_risk),
            "still_waiting": len(waiting),
        },
        "itc_claimable": claimable,
        "itc_at_risk": sum((_tax(inv) for inv in at_risk), Decimal(0)),
        "rows": rows,
        "books_only": at_risk,
        "still_waiting": waiting,
    }


# ---------------------------------------------------------------------------
# A sample statement for the demo (clearly synthetic)
# ---------------------------------------------------------------------------

SAMPLE_FORMAT = "synthetic sample (from your books)"


def sample_for(db: Session, user_id: int, period: date) -> dict:
    """A synthetic GSTR-2B for `period`, built from this user's approved invoices so the
    Reconcile page can be tried without a portal file. To show every kind of result:
    the 2nd invoice is reported with a higher amount, the last is left out (ITC at
    risk) and one invoice from an unbooked supplier is added."""
    end = _period_end(period)
    books = db.scalars(
        select(Invoice)
        .where(
            Invoice.user_id == user_id,
            Invoice.status == InvoiceStatus.APPROVED,
            Invoice.invoice_date >= period,
            Invoice.invoice_date <= end,
            Invoice.vendor_gstin.is_not(None),
            Invoice.invoice_number.is_not(None),
        )
        .order_by(Invoice.invoice_date, Invoice.id)
    ).all()
    suppliers: dict[str, dict] = {}
    for position, invoice in enumerate(books):
        if len(books) >= 3 and position == len(books) - 1:
            break  # left out: shows "ITC at risk"
        taxable = _money(invoice.subtotal)
        taxes = {k: _money(getattr(invoice, k)) for k in ("igst", "cgst", "sgst")}
        value = _money(invoice.total)
        if len(books) >= 2 and position == 1:  # reported higher: shows "mismatched"
            taxable += 100
            value += 100
        supplier = suppliers.setdefault(
            invoice.vendor_gstin,
            {"ctin": invoice.vendor_gstin, "trdnm": invoice.vendor_name, "inv": []},
        )
        supplier["inv"].append(
            {
                "inum": invoice.invoice_number,
                "dt": invoice.invoice_date.strftime("%d-%m-%Y"),
                "val": str(value),
                "rev": "N",
                "itcavl": "Y",
                "items": [{"txval": str(taxable), **{k: str(v) for k, v in taxes.items()}}],
            }
        )
    suppliers["24SAMPL0000A1Z5"] = {
        "ctin": "24SAMPL0000A1Z5",
        "trdnm": "Sample Unbooked Supplier",
        "inv": [
            {
                "inum": "SAMPLE-1",
                "dt": (period + timedelta(days=9)).strftime("%d-%m-%Y"),
                "val": "1180.00",
                "rev": "N",
                "itcavl": "Y",
                "items": [{"txval": "1000.00", "igst": "180.00"}],
            }
        ],
    }
    return {
        "data": {
            "rtnprd": period.strftime("%m%Y"),
            "version": "synthetic-sample",
            "docdata": {"b2b": list(suppliers.values())},
        }
    }
