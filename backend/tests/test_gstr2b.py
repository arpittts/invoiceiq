"""GSTR-2B: parsing the documented layout and matching it against the approved books."""

import hashlib
import json
from datetime import date, datetime
from decimal import Decimal as D

import pytest
from sqlalchemy import select

from app.models import Invoice, InvoiceFile, InvoiceStatus, MatchStatus, User
from app.security import hash_password
from app.services import gstr2b
from app.services.gstr2b import Gstr2bFormatError, number_key, parse
from tests.conftest import register_and_login

A = "27AAPFU0939F1ZV"  # supplier A
B = "29AABCU9603R1ZM"  # supplier B


def two_b(*suppliers, period="072026", wrap=True):
    """A GSTR-2B JSON in the documented layout. suppliers: (ctin, [inv dicts])."""
    body = {
        "gstin": "27AAACR5055K1Z7",
        "rtnprd": period,
        "docdata": {
            "b2b": [{"ctin": c, "trdnm": f"Supplier {c[:2]}", "inv": inv} for c, inv in suppliers]
        },
    }
    return {"data": body} if wrap else body


def inv(number, day, taxable, igst="0", cgst="0", sgst="0", value=None, itcavl="Y"):
    total = D(taxable) + D(igst) + D(cgst) + D(sgst)
    return {
        "inum": number,
        "dt": day,
        "val": str(value or total),
        "pos": "27",
        "rev": "N",
        "itcavl": itcavl,
        "items": [
            {"rt": 18, "txval": taxable, "igst": igst, "cgst": cgst, "sgst": sgst, "cess": 0}
        ],
    }


@pytest.fixture
def user(db):
    u = User(email="books@example.com", password_hash=hash_password("not-used-123"))
    db.add(u)
    db.commit()
    return u


def book(db, user, gstin, number, day, taxable, igst="0", cgst="0", sgst="0"):
    taxes = D(igst) + D(cgst) + D(sgst)
    invoice = Invoice(
        user_id=user.id,
        status=InvoiceStatus.APPROVED,
        vendor_gstin=gstin,
        invoice_number=number,
        invoice_date=day,
        subtotal=D(taxable),
        igst=D(igst) if D(igst) else None,
        cgst=D(cgst) if D(cgst) else None,
        sgst=D(sgst) if D(sgst) else None,
        total=D(taxable) + taxes,
        file=InvoiceFile(
            user_id=user.id,
            original_name="x.pdf",
            mime_type="application/pdf",
            storage_path="-",
            page_count=1,
            sha256=hashlib.sha256(f"{gstin}{number}".encode()).hexdigest(),
        ),
    )
    db.add(invoice)
    db.commit()
    return invoice


def upload(db, user, data, name="2b.json"):
    statement = gstr2b.save_statement(db, user.id, parse(data), file_name=name)
    db.commit()
    return statement


def by_number(statement):
    return {row.invoice_number: row for row in statement.invoices}


# --- parsing ----------------------------------------------------------------------------


def test_parses_the_documented_layout():
    s = parse(two_b((A, [inv("INV-1", "05-07-2026", "1000.00", igst="180.00")])))
    assert s.period == date(2026, 7, 1) and s.recipient_gstin == "27AAACR5055K1Z7"
    [row] = s.rows
    assert (row.supplier_gstin, row.invoice_number, row.invoice_date) == (
        A,
        "INV-1",
        date(2026, 7, 5),
    )
    assert (row.taxable_value, row.igst, row.invoice_value) == (
        D("1000.00"),
        D("180.00"),
        D("1180.00"),
    )
    assert row.itc_available and not row.reverse_charge


def test_accepts_layout_variations():
    """No outer 'data', items called 'itms' with amounts nested in 'itm_det'."""
    data = two_b(wrap=False)
    data["docdata"]["b2b"] = [
        {
            "ctin": A,
            "inv": [
                {
                    "inum": "9",
                    "dt": "01-07-2026",
                    "val": 590,
                    "itms": [{"itm_det": {"txval": 500, "camt": 45, "samt": 45}}],
                }
            ],
        }
    ]
    [row] = parse(data).rows
    assert (row.cgst, row.sgst, row.taxable_value) == (D("45.00"), D("45.00"), D("500.00"))


@pytest.mark.parametrize(
    ("data", "message"),
    [
        ([], "JSON object"),
        ({"rtnprd": "132026", "docdata": {}}, "no month 13"),
        ({"rtnprd": "072026"}, "docdata"),
        ({"rtnprd": "072026", "docdata": {"b2b": [{"inv": []}]}}, "ctin"),
        (two_b((A, [{**inv("1", "05-07-2026", "1"), "dt": "July 5"}])), r"inv\[0\].dt"),
        (two_b((A, [{**inv("1", "05-07-2026", "1"), "val": "abc"}])), r"inv\[0\].val"),
    ],
)
def test_explains_files_it_cannot_read(data, message):
    with pytest.raises(Gstr2bFormatError, match=message):
        parse(data)


@pytest.mark.parametrize(
    ("a", "b"),
    [("INV/26-27/0042", "inv 26 27 42"), ("0042", "42"), ("SH-0001", "sh/1"), ("A1B2", "a1 b2")],
)
def test_invoice_numbers_match_despite_formatting(a, b):
    assert number_key(a) == number_key(b)


def test_different_numbers_stay_different():
    assert number_key("INV-42") != number_key("INV-43")
    assert number_key("2026-27/42") != number_key("2025-26/42")


# --- matching ---------------------------------------------------------------------------


def test_each_kind_of_result(db, user):
    book(db, user, A, "INV-1", date(2026, 7, 5), "1000.00", igst="180.00")  # matches
    book(db, user, A, "INV-2", date(2026, 7, 6), "2000.00", igst="360.00")  # amount differs
    book(db, user, B, "B/26-27/0007", date(2026, 7, 9), "500.00", cgst="45", sgst="45")  # number
    book(db, user, B, "B-9", date(2026, 7, 20), "800.00", cgst="72", sgst="72")  # not reported
    statement = upload(
        db,
        user,
        two_b(
            (
                A,
                [
                    inv("INV-1", "05-07-2026", "1000.00", igst="180.00"),
                    inv("INV-2", "06-07-2026", "2100.00", igst="378.00"),
                    inv("INV-3", "07-07-2026", "300.00", igst="54.00"),  # not in books
                ],
            ),
            (B, [inv("7", "09-07-2026", "500.00", cgst="45", sgst="45")]),
        ),
    )
    rows = by_number(statement)
    assert rows["INV-1"].match_status == MatchStatus.MATCHED
    assert rows["INV-2"].match_status == MatchStatus.MISMATCHED
    assert {d["field"] for d in rows["INV-2"].differences} == {"taxable_value", "igst", "total"}
    assert rows["INV-3"].match_status == MatchStatus.ONLY_2B
    assert rows["7"].match_status == MatchStatus.POSSIBLE
    assert rows["7"].differences[0]["books"] == "B/26-27/0007"

    result = gstr2b.report(db, statement)
    assert result["counts"] == {
        "matched": 1, "mismatched": 1, "possible": 1, "only_2b": 1,
        "books_only": 1, "still_waiting": 0,
    }  # fmt: skip
    assert [inv.invoice_number for inv in result["books_only"]] == ["B-9"]
    assert result["itc_at_risk"] == D("144.00")  # B-9's CGST + SGST
    assert result["itc_claimable"] == D("180.00")  # only the clean match


def test_reformatted_number_of_the_same_invoice_is_an_exact_match(db, user):
    book(db, user, A, "INV/26-27/0042", date(2026, 7, 5), "1000.00", igst="180.00")
    statement = upload(
        db, user, two_b((A, [inv("inv 26 27 42", "05-07-2026", "1000", igst="180")]))
    )
    assert statement.invoices[0].match_status == MatchStatus.MATCHED


def test_rounding_within_one_rupee_still_matches(db, user):
    book(db, user, A, "1", date(2026, 7, 5), "1000.00", igst="180.00")
    statement = upload(
        db, user, two_b((A, [inv("1", "05-07-2026", "1000.40", igst="180.07", value="1181")]))
    )
    assert statement.invoices[0].match_status == MatchStatus.MATCHED


def test_a_late_reported_invoice_matches_next_month(db, user):
    """A June bill reported in July's 2B: matched, and not 'at risk' in July."""
    book(db, user, A, "JUNE-1", date(2026, 6, 28), "1000.00", igst="180.00")
    june = upload(db, user, two_b(period="062026"))
    assert gstr2b.report(db, june)["counts"]["books_only"] == 1  # at risk in June
    july = upload(db, user, two_b((A, [inv("JUNE-1", "28-06-2026", "1000", igst="180")])))
    assert july.invoices[0].match_status == MatchStatus.MATCHED
    assert gstr2b.report(db, july)["counts"]["still_waiting"] == 0


def test_an_older_unmatched_invoice_is_still_waiting_not_at_risk(db, user):
    book(db, user, A, "MAY-1", date(2026, 5, 3), "100.00", igst="18.00")
    result = gstr2b.report(db, upload(db, user, two_b()))
    assert result["counts"]["books_only"] == 0 and result["counts"]["still_waiting"] == 1


def test_an_invoice_matched_by_one_month_is_not_matched_again(db, user):
    book(db, user, A, "INV-1", date(2026, 6, 5), "1000.00", igst="180.00")
    upload(db, user, two_b((A, [inv("INV-1", "05-06-2026", "1000", igst="180")]), period="062026"))
    july = upload(db, user, two_b((A, [inv("INV-1", "05-06-2026", "1000", igst="180")])))
    assert july.invoices[0].match_status == MatchStatus.ONLY_2B  # a duplicate report


def test_only_approved_invoices_of_this_user_count(db, user):
    other = User(email="other@example.com", password_hash=hash_password("not-used-123"))
    db.add(other)
    db.commit()
    book(db, other, A, "INV-1", date(2026, 7, 5), "1000.00", igst="180.00")
    pending = book(db, user, A, "INV-2", date(2026, 7, 5), "1000.00", igst="180.00")
    pending.status = InvoiceStatus.NEEDS_REVIEW
    db.commit()
    statement = upload(
        db,
        user,
        two_b(
            (
                A,
                [
                    inv("INV-1", "05-07-2026", "1000", igst="180"),
                    inv("INV-2", "05-07-2026", "1000", igst="180"),
                ],
            )
        ),
    )
    assert {r.match_status for r in statement.invoices} == {MatchStatus.ONLY_2B}


def test_uploading_a_month_again_replaces_it(db, user):
    upload(db, user, two_b((A, [inv("1", "05-07-2026", "100")])), name="first.json")
    second = upload(db, user, two_b((A, [inv("2", "05-07-2026", "100")])), name="second.json")
    assert second.file_name == "second.json"
    assert [r.invoice_number for r in second.invoices] == ["2"]
    from sqlalchemy import func

    from app.models import Gstr2bStatement

    assert db.scalar(select(func.count()).select_from(Gstr2bStatement)) == 1


def test_itc_not_available_is_not_counted_as_claimable(db, user):
    book(db, user, A, "1", date(2026, 7, 5), "1000.00", igst="180.00")
    statement = upload(
        db, user, two_b((A, [inv("1", "05-07-2026", "1000", igst="180", itcavl="N")]))
    )
    assert statement.invoices[0].match_status == MatchStatus.MATCHED
    assert gstr2b.report(db, statement)["itc_claimable"] == D(0)


# --- API ---------------------------------------------------------------------------------


def upload_file(client, headers, data, name="2b.json"):
    return client.post(
        "/api/v1/reconcile/gstr2b",
        headers=headers,
        files={"file": (name, json.dumps(data).encode(), "application/json")},
    )


def test_api_upload_report_recheck_and_csv(client, auth_headers, db):
    from app.models import User as UserModel

    owner = db.scalar(select(UserModel))
    book(db, owner, A, "INV-1", date(2026, 7, 5), "1000.00", igst="180.00")
    data = two_b(
        (A, [inv("INV-1", "05-07-2026", "1000", igst="180"), inv("INV-2", "06-07-2026", "50")])
    )
    response = upload_file(client, auth_headers, data)
    assert response.status_code == 201
    body = response.json()
    assert body["period"] == "2026-07" and body["verified_against_portal"] is False  # JSON
    assert body["counts"]["matched"] == 1 and body["counts"]["only_2b"] == 1
    assert client.get("/api/v1/reconcile/gstr2b", headers=auth_headers).json()[0]["period"] == (
        "2026-07"
    )

    # The supplier's second invoice gets approved later: re-check matches it.
    book(db, owner, A, "INV-2", date(2026, 7, 6), "50.00")
    rechecked = client.post("/api/v1/reconcile/gstr2b/2026-07/recheck", headers=auth_headers)
    assert rechecked.json()["counts"]["matched"] == 2

    csv_response = client.get("/api/v1/reconcile/gstr2b/2026-07/export.csv", headers=auth_headers)
    assert csv_response.status_code == 200
    assert csv_response.text.splitlines()[0].startswith("result,supplier_gstin")
    assert "attachment" in csv_response.headers["content-disposition"]


def test_api_explains_unreadable_files(client, auth_headers):
    bad_json = client.post(
        "/api/v1/reconcile/gstr2b",
        headers=auth_headers,
        files={"file": ("x.json", b"{not json", "application/json")},
    )
    assert bad_json.status_code == 422 and "Excel" in bad_json.json()["detail"]
    not_2b = upload_file(client, auth_headers, {"hello": "world"})
    assert not_2b.status_code == 422 and "GSTR-2B" in not_2b.json()["detail"]


def test_api_reports_are_private(client, auth_headers):
    upload_file(client, auth_headers, two_b())
    stranger = register_and_login(client, "stranger@example.com")
    assert client.get("/api/v1/reconcile/gstr2b/2026-07", headers=stranger).status_code == 404
    assert client.get("/api/v1/reconcile/gstr2b", headers=stranger).json() == []
    assert client.get("/api/v1/reconcile/gstr2b/2026-13x", headers=auth_headers).status_code == 422


def test_csv_defuses_formulas_from_supplier_text(client, auth_headers):
    data = two_b((A, [inv("=CMD()", "05-07-2026", "100")]))
    upload_file(client, auth_headers, data)
    text = client.get("/api/v1/reconcile/gstr2b/2026-07/export.csv", headers=auth_headers).text
    assert "'=CMD()" in text


def test_a_sample_statement_shows_every_kind_of_result(client, auth_headers, db):
    from app.models import User as UserModel

    owner = db.scalar(select(UserModel))
    for n, day in enumerate([3, 9, 17], start=1):
        book(db, owner, A, f"S-{n}", date(2026, 7, day), "1000.00", igst="180.00")
    response = client.post("/api/v1/reconcile/gstr2b/sample/2026-07", headers=auth_headers)
    assert response.status_code == 201
    body = response.json()
    assert body["source_format"].startswith("synthetic sample")
    assert body["counts"] == {
        "matched": 1, "mismatched": 1, "possible": 0, "only_2b": 1,
        "books_only": 1, "still_waiting": 0,
    }  # fmt: skip
    assert [b["invoice_number"] for b in body["books_only"]] == ["S-3"]
    assert all(len(row["supplier_gstin"]) == 15 for row in body["rows"])  # never truncated


# --- the portal's Excel download (layout copied from a real file; values made up) --------

EXCEL_TOP = [
    "GSTIN of supplier", "Trade/Legal name", "Invoice Details", "", "", "", "Place of supply",
    "Supply Attract Reverse Charge", "Taxable Value (₹)", "Tax Amount", "", "", "",
    "GSTR-1/1A/IFF/GSTR-5 Period", "GSTR-1/1A/IFF/GSTR-5 Filing Date", "ITC Availability",
    "Reason", "Applicable % of Tax Rate", "Source", "IRN", "IRN Date",
]  # fmt: skip
EXCEL_SUB = [
    "", "", "Invoice number", "Invoice type", "Invoice Date", "Invoice Value(₹)", "", "", "",
    "Integrated Tax(₹)", "Central Tax(₹)", "State/UT Tax(₹)", "Cess(₹)", "", "", "", "", "", "",
    "", "",
]  # fmt: skip


def excel_2b(rows, month="July", year="2026-27") -> bytes:
    """A workbook shaped like the portal's GSTR-2B download, with made-up values."""
    import io

    import openpyxl

    wb = openpyxl.Workbook()
    readme = wb.active
    readme.title = "Read me"
    readme.append(["Goods and Services Tax  - GSTR-2B"])
    readme.append([])
    readme.append([])
    for label, value in [("Financial Year", year), ("Tax Period", month),
                         ("GSTIN", "27AAACR5055K1Z7"), ("Legal Name", "Test Buyer")]:  # fmt: skip
        readme.append([label, value])
    b2b = wb.create_sheet("B2B")
    b2b.append(["Goods and Services Tax  - GSTR-2B"])
    b2b.append([])
    b2b.append([])
    b2b.append(["Taxable inward supplies received from registered persons"])
    b2b.append(EXCEL_TOP)
    b2b.append(EXCEL_SUB)
    for row in rows:
        b2b.append(row)
    b2b.append([])  # trailing blank row, as in a real download
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def excel_row(gstin, number, day, value, place, taxable, igst=0, cgst=0, sgst=0, itc="Yes"):
    return [gstin, "Made Up Traders", number, "Regular", day, value, place, "No", taxable,
            igst, cgst, sgst, 0, "Jul'26", "11/08/2026", itc, "", "100%", "", "", ""]  # fmt: skip


def test_reads_the_portal_excel_layout():
    content = excel_2b(
        [
            excel_row(
                A,
                "RA/2026-27/0001",
                "11/07/2026",
                10597,
                "Maharashtra",
                8980.8,
                cgst=808.27,
                sgst=808.27,
            ),  # fmt: skip
            excel_row(B, "77", datetime(2026, 7, 2), 1180, "Karnataka", 1000, igst=180, itc="No"),
        ]
    )
    s = gstr2b.parse_file(content, "2b.xlsx")
    assert s.source_format == gstr2b.EXCEL_FORMAT and s.period == date(2026, 7, 1)
    assert s.recipient_gstin == "27AAACR5055K1Z7"
    first, second = s.rows
    assert (first.supplier_gstin, first.invoice_number, first.invoice_date) == (
        A, "RA/2026-27/0001", date(2026, 7, 11),
    )  # fmt: skip
    assert (first.taxable_value, first.cgst, first.sgst) == (D("8980.80"), D("808.27"), D("808.27"))
    assert first.place_of_supply == "27" and not first.reverse_charge and first.itc_available
    assert second.invoice_date == date(2026, 7, 2)  # a real date cell works too
    assert second.place_of_supply == "29" and second.igst == D("180.00")
    assert not second.itc_available


@pytest.mark.parametrize(
    ("month", "year", "period"),
    [("July", "2026-27", date(2026, 7, 1)), ("January", "2026-27", date(2027, 1, 1)),
     ("March", "2025-26", date(2026, 3, 1)), ("April", "2026-27", date(2026, 4, 1))],
)  # fmt: skip
def test_excel_period_follows_the_april_to_march_financial_year(month, year, period):
    assert gstr2b.parse_file(excel_2b([], month, year), "2b.xlsx").period == period


def test_explains_excel_files_it_cannot_read():
    import io

    import openpyxl

    wb = openpyxl.Workbook()
    buffer = io.BytesIO()
    wb.save(buffer)
    with pytest.raises(Gstr2bFormatError, match="Read me"):
        gstr2b.parse_file(buffer.getvalue(), "x.xlsx")
    with pytest.raises(Gstr2bFormatError, match="Excel"):
        gstr2b.parse_file(b"PK\x03\x04 not really a zip", "x.xlsx")
    with pytest.raises(Gstr2bFormatError, match="Excel .* or JSON"):
        gstr2b.parse_file(b"\x00\x01binary", "x.bin")


def test_api_accepts_the_portal_excel_and_marks_it_verified(client, auth_headers, db):
    from app.models import User as UserModel

    owner = db.scalar(select(UserModel))
    book(db, owner, A, "INV-1", date(2026, 7, 5), "1000.00", igst="180.00")
    content = excel_2b([excel_row(A, "inv 1", "05/07/2026", 1180, "Maharashtra", 1000, igst=180)])
    response = client.post(
        "/api/v1/reconcile/gstr2b",
        headers=auth_headers,
        files={"file": ("072026_GSTR2B.xlsx", content, "application/octet-stream")},
    )
    assert response.status_code == 201
    body = response.json()
    assert body["verified_against_portal"] is True
    assert body["source_format"] == gstr2b.EXCEL_FORMAT
    assert body["counts"]["matched"] == 1
