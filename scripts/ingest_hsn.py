"""Load the official HSN/SAC code list and indicative GST rates into hsn_codes.

Inputs (downloaded by hand, kept out of git):
- HSN_SAC.xlsx: the GST portal's HSN directory (sheets HSN_MSTR and SAC_MSTR)
- Notification 09/2025-Central Tax (Rate): goods rate schedules (CGST %; GST = 2x)
- Notification 01/2026-Central Tax (Rate): its May 2026 amendment (applied below)

One row per code, no text splitting: each code is already a small, self-contained
document. Each row gets its parent heading as context, an embedding of
"code description. heading", and the rate of its most specific schedule entry.

Rates are INDICATIVE: many schedule entries carry conditions in their description
("pre-packaged and labelled", "of value not exceeding Rs 2500"), and exemptions live
in other notifications. Each row's source_note names the schedule entry it used,
so a reviewer can check the condition. Amendment 19/2025 (31 Dec 2025) is not applied.

Re-running is safe: existing codes are updated in place (ids stay stable, so saved
suggestions stay linked), new codes are added and codes no longer listed removed.

Usage (from the repo root; DATABASE_URL picks the database):
    backend/.venv/bin/python scripts/ingest_hsn.py --dir "data/hsn raw"
"""

import argparse
import re
import sys
import time
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "backend"))

DIRECTORY_FILE = "HSN_SAC.xlsx"
RATE_NOTIFICATION = "266209.pdf"  # 09/2025-Central Tax (Rate), as published in the Gazette
FIRST_ENGLISH_PAGE = 76  # the Gazette prints Hindi first; English starts on page 76

# Notification 01/2026-Central Tax (Rate): column (2) of these entries is replaced.
AMENDMENT_01_2026 = {
    ("I", 150): "2202 99 21, 2202 99 29",
    ("I", 151): "22029931, 22029939",
    ("III", 2): "2202 91 00, 2202 99 91, 2202 99 99",
    ("III", 3): "2202 99 91, 2202 99 99",
}


# ---------------------------------------------------------------------------
# Parsing the code column of a schedule
# ---------------------------------------------------------------------------


def parse_codes(text: str) -> list[str]:
    """'2202 99 21, 2202 99 29' -> ['22029921', '22029929']. Also handles ranges
    ('3901 to 3913'), split digit groups ('27 01') and ignores exclusions in brackets
    ('9603 [other than 9603 10 00]') and words ('Any Chapter')."""
    text = re.sub(r"\[[^\]]*\]|\([^)]*\)", " ", text)  # exclusions and notes
    codes = []
    for segment in re.split(r",|;|\bor\b|\band\b", text):
        segment = segment.strip()
        if " to " in segment:
            start, _, end = segment.partition(" to ")
            a, b = re.sub(r"\D", "", start), re.sub(r"\D", "", end)
            if a and b and len(a) == len(b) and len(a) <= 4 and int(a) <= int(b):
                codes += [str(n).zfill(len(a)) for n in range(int(a), int(b) + 1)]
            continue
        digits = "".join(re.findall(r"\d+", segment))
        if len(digits) in (2, 4, 6, 8):
            codes.append(digits)
    return codes


@dataclass
class RateEntry:
    schedule: str  # "I", "II", ...
    serial: int
    codes_text: str
    description: str
    gst_rate: Decimal  # total GST: twice the CGST rate in the schedule
    codes: list[str] = field(default_factory=list)


def gst_from_heading(heading: str) -> tuple[str, Decimal] | None:
    """'Schedule II – 9 %' -> ('II', Decimal('18'))."""
    match = re.search(r"Schedule\s*[–-]?\s*([IVX]+)\s*[–-]\s*([\d.]+)\s*%", heading)
    if not match:
        return None
    return match.group(1), Decimal(match.group(2)) * 2


def parse_rate_schedules(pdf_path: Path) -> list[RateEntry]:
    import pymupdf

    entries: list[RateEntry] = []
    schedule: tuple[str, Decimal] | None = None
    doc = pymupdf.open(pdf_path)
    for page in doc.pages(FIRST_ENGLISH_PAGE - 1, doc.page_count):
        # Schedule headings and tables, top to bottom, so each table row gets the
        # schedule whose heading came last above it.
        items = []
        for block in page.get_text("blocks"):
            found = gst_from_heading(" ".join(block[4].split()))
            if found:
                items.append((block[1], "heading", found))
        for table in page.find_tables().tables:
            items.append((table.bbox[1], "table", table))
        for _, kind, item in sorted(items, key=lambda entry: entry[0]):
            if kind == "heading":
                schedule = item
                continue
            if schedule is None:
                continue
            for row in item.extract():
                cells = [" ".join((cell or "").split()) for cell in row]
                if len(cells) < 3:
                    continue
                serial, codes_text, description = cells[0], " ".join(cells[1:-1]), cells[-1]
                number = re.fullmatch(r"(\d+)[A-Z]?\.?", serial)
                if number:
                    entries.append(
                        RateEntry(
                            schedule[0], int(number.group(1)), codes_text, description, schedule[1]
                        )
                    )
                elif entries and not serial:  # a row continued from the previous one
                    entries[-1].codes_text = f"{entries[-1].codes_text} {codes_text}".strip()
                    entries[-1].description = f"{entries[-1].description} {description}".strip()

    for entry in entries:
        entry.codes_text = AMENDMENT_01_2026.get((entry.schedule, entry.serial), entry.codes_text)
        entry.codes = parse_codes(entry.codes_text)
    return entries


# ---------------------------------------------------------------------------
# Assigning a rate to each directory code
# ---------------------------------------------------------------------------


def pct(rate: Decimal) -> str:
    """Decimal('40.00') -> '40' (normalize() alone would give '4E+1')."""
    return format(rate.normalize(), "f")


def assign_rate(code: str, by_prefix: dict[str, list[RateEntry]]) -> tuple[Decimal | None, str]:
    """The rate of the most specific schedule entry covering `code`, plus a note.

    If entries at that level disagree (the rate depends on the description), no
    single rate is given; the note lists the possible rates instead.
    """
    for length in range(len(code), 1, -1):
        matches = by_prefix.get(code[:length])
        if not matches:
            continue
        rates = sorted({entry.gst_rate for entry in matches})
        cited = "; ".join(
            f"Sch {e.schedule} S.No {e.serial} ({pct(e.gst_rate)}%): {e.description[:120]}"
            for e in matches[:3]
        )
        if len(rates) == 1:
            return rates[0], cited
        options = " or ".join(f"{pct(rate)}%" for rate in rates)
        return None, f"Rate depends on the description ({options}). {cited}"
    return None, "No entry in Notification 09/2025 (may be exempt or listed elsewhere)"


# ---------------------------------------------------------------------------
# Reading the directory
# ---------------------------------------------------------------------------


@dataclass
class CodeRow:
    code: str
    kind: str
    description: str
    heading: str | None = None


def read_directory(path: Path) -> list[CodeRow]:
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True)
    rows: dict[str, CodeRow] = {}
    for sheet, kind in (("HSN_MSTR", "HSN"), ("SAC_MSTR", "SAC")):
        for i, (raw_code, raw_description) in enumerate(
            workbook[sheet].iter_rows(values_only=True, max_col=2)
        ):
            if i == 0 or raw_code is None or not raw_description:
                continue
            code = re.sub(r"\D", "", str(raw_code))
            if kind == "HSN" and len(code) % 2:  # Excel may drop a leading zero: 101 -> 0101
                code = "0" + code
            if len(code) < 2 or code in rows:
                continue
            rows[code] = CodeRow(code, kind, " ".join(str(raw_description).split()))

    # Context for each code: the text of its parent headings (4- and 6-digit).
    for row in rows.values():
        parents = [
            rows[row.code[:n]].description
            for n in (4, 6)
            if n < len(row.code) and row.code[:n] in rows
        ]
        row.heading = " / ".join(parents) or None
    return list(rows.values())


def is_fragment(description: str) -> bool:
    """Rows like 'OTHER', 'FOR TWO WHEELERS' or 'IN SOLID FORM' need their heading."""
    words = re.findall(r"[A-Za-z]+", description)
    return len(words) <= 3 or bool(re.match(r"(?i)\W*(others?|for|of|in|with|not)\b", description))


def embedding_text(code: str, description: str, heading: str | None) -> str:
    """What gets embedded. Fragments ('OTHER', 'FOR TWO WHEELERS') lead with their
    heading, so 'OTHER' under 'ALUMINIUM WIRE' is searched as aluminium wire.

    Measured on the labelled set with a freshly rebuilt index: hybrid heading
    recall@1 63.0% -> 68.0% (vector-only alone dipped slightly, 72.4% -> 71.8% @5).
    """
    if heading and is_fragment(description):
        return f"{heading}: {description.lower()} (code {code})"
    return f"{code} {description}. {heading or ''}".strip()


# ---------------------------------------------------------------------------


def main() -> None:
    parser = argparse.ArgumentParser(description="Load HSN/SAC codes and rates into hsn_codes")
    parser.add_argument("--dir", default="data/hsn raw", help="folder with the downloaded files")
    parser.add_argument(
        "--downloaded",
        default=date.today().isoformat(),
        help="date the files were downloaded (for source notes)",
    )
    parser.add_argument("--dry-run", action="store_true", help="parse and report, don't write")
    parser.add_argument(
        "--embed",
        choices=["all", "fragments"],
        default="all",
        help="recompute embeddings for all rows, or only fragment rows",
    )
    args = parser.parse_args()
    folder = (REPO / args.dir) if not Path(args.dir).is_absolute() else Path(args.dir)

    started = time.time()
    codes = read_directory(folder / DIRECTORY_FILE)
    entries = parse_rate_schedules(folder / RATE_NOTIFICATION)
    by_prefix: dict[str, list[RateEntry]] = {}
    for entry in entries:
        for code in entry.codes:
            by_prefix.setdefault(code, []).append(entry)

    source = f"GST portal HSN directory (downloaded {args.downloaded})"
    records = []
    for row in codes:
        rate, rate_note = (None, "Services rates not loaded (no consolidated notification)")
        if row.kind == "HSN":
            rate, rate_note = assign_rate(row.code, by_prefix)
        records.append(
            {
                "code": row.code,
                "kind": row.kind,
                "description": row.description,
                "heading": row.heading,
                "chapter": row.code[:2],
                "default_gst_rate": rate,
                "source_note": f"{source}. Rate (indicative, Notification 09/2025-CT(Rate) "
                f"as amended by 01/2026): {rate_note}",
            }
        )

    hsn = [r for r in records if r["kind"] == "HSN"]
    with_rate = sum(r["default_gst_rate"] is not None for r in hsn)
    ambiguous = sum("depends on the description" in r["source_note"] for r in hsn)
    print(f"Directory: {len(hsn)} HSN + {len(records) - len(hsn)} SAC codes")
    print(
        f"Rate schedule entries parsed: {len(entries)} "
        f"({sum(not e.codes for e in entries)} without a code, e.g. 'Any Chapter')"
    )
    print(
        f"HSN codes with a rate: {with_rate} ({with_rate / len(hsn):.0%}); "
        f"rate depends on description: {ambiguous}; no entry: {len(hsn) - with_rate - ambiguous}"
    )
    if args.dry_run:
        return

    from sqlalchemy import select

    from app.db import SessionLocal
    from app.llm.embeddings import get_embedder
    from app.models import HsnCode

    embedder = get_embedder()
    to_embed = [
        i for i, r in enumerate(records) if args.embed == "all" or is_fragment(r["description"])
    ]
    texts = [
        embedding_text(records[i]["code"], records[i]["description"], records[i]["heading"])
        for i in to_embed
    ]
    vectors: dict[int, list[float]] = {}
    for start in range(0, len(texts), 64):  # small batches keep memory modest on long texts
        batch = embedder.embed_documents(texts[start : start + 64])
        vectors.update(zip(to_embed[start : start + 64], batch, strict=True))
        print(f"\rEmbedding {min(start + 64, len(texts))}/{len(texts)}", end="", flush=True)
    print()

    db = SessionLocal()
    try:
        existing = {row.code: row for row in db.scalars(select(HsnCode))}
        seen = set()
        for i, record in enumerate(records):
            row = existing.get(record["code"]) or HsnCode(code=record["code"])
            for name, value in record.items():
                setattr(row, name, value)
            if i in vectors:
                row.embedding = vectors[i]
            db.add(row)
            seen.add(record["code"])
        removed = [row for code, row in existing.items() if code not in seen]
        for row in removed:
            db.delete(row)  # line items citing it keep their code text; the link becomes NULL
        db.commit()
        # Updates leave dead rows behind, and pgvector's HNSW index then returns
        # fewer live candidates (we measured vector recall@5 falling from 72% to 62%).
        # Clean up and rebuild the index so search quality doesn't silently degrade.
        with db.get_bind().connect().execution_options(isolation_level="AUTOCOMMIT") as conn:
            conn.exec_driver_sql("VACUUM ANALYZE hsn_codes")
            conn.exec_driver_sql("REINDEX INDEX ix_hsn_codes_embedding")
        print(
            f"Saved {len(seen)} codes ({len(seen) - len(set(existing) & seen)} new, "
            f"{len(removed)} removed) in {time.time() - started:.0f}s"
        )
    finally:
        db.close()


if __name__ == "__main__":
    main()
